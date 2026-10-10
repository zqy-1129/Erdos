"""Actual local calculation and draft assembly. The task classifier is an explicit fixture."""
import argparse,contextlib,csv,hashlib,importlib.metadata,io,sys,uuid
from pathlib import Path
from types import SimpleNamespace
import content_runtime as rt
import stage03_fileio as fio
from consumer_v2 import ReleaseStore,DEFAULT_RELEASE,startup,cmd_outline,cmd_search
def call(fn,store,args):
 buf=io.StringIO()
 with contextlib.redirect_stdout(buf):code=fn(store,args)
 if code not in [0,1]:raise ValueError('CONSUMER_FAILED')
 return code,rt.parse(buf.getvalue())
def tex_escape(s):
 return ''.join({'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}.get(c,c) for c in s)
def narrative(section,results):
 q=section['subproblem_index'];kind=section['section_type']
 if kind=='abstract':return '合成示例，采用最小二乘拟合和网格搜索；仅用于检验当前任务的数据、计算与文稿链路。'
 if kind=='restatement':return '问题一估计合成运量与单位成本的线性关系；问题二独立最小化输入配置中定义的二次成本。两问为演示任务。'
 if kind=='assumptions':return '合成样本用于回归示例；二次成本函数与搜索范围为显式输入，不是从历史论文取得。'
 if q==1 and kind in ['model','solve','analysis','validation']:
  r=results['q1']
  if kind=='analysis':return '使用本次保存并重新读取的 CSV 中 %d 个合成样本，检查运量与单位成本的线性关系；未使用历史论文数据。'%r['n_samples']
  if kind=='model':return '建立 y=a·v+b 的线性模型，以样本残差平方和最小作为拟合目标；a 表示斜率，b 表示截距。'
  if kind=='validation':return '训练样本上的 R² 为 %.6f、RMSE 为 %.6f，残差图用于检查拟合表现；本示例未做独立测试集验证，不能据此宣称泛化能力。'%(r['r_squared'],r['rmse'])
  return '最小二乘计算得到截距 %.6f、斜率 %.6f。以下拟合与残差图均使用本次 CSV 生成。'%(r['intercept'],r['slope'])
 if q==2 and kind in ['model','solve','analysis','validation']:
  r=results['q2']
  if kind=='analysis':return '在运量区间 [1,10] 内最小化独立输入的成本函数。它是本次演示给定的假设，未由问题一拟合结果或历史论文推导。'
  if kind=='model':return '目标函数为 C(v)=(v−%.2f)²+%.2f，约束为 1≤v≤10；凸二次函数存在唯一的区间内极小值。'%(r['center'],r['offset'])
  if kind=='validation':return '导数为零的解析最优点 %.6f，与网格解相差 %.6f；网格只近似连续最优点，二者误差已记录。'%(r['analytic_volume'],r['grid_error'])
  return '在 %d 个网格点中搜索得到 v=%.6f、成本 %.8f；以下目标曲线来自当前 objective.json。'%(r['grid_points'],r['opt_volume'],r['opt_cost'])
 if kind=='symbols':return 'v：运量（示例单位）；y：单位配送成本（示例单位）；C：独立总成本（示例单位）；a、b：线性模型系数。'
 return '待填写：本示例仅完成两问的局部计算，该章节尚无独立证据或完整论述。'
def main(argv=None):
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--release',default=DEFAULT_RELEASE);p.add_argument('--out',default='out/demos');p.add_argument('--run-id');a=p.parse_args(argv)
 rid=rt.identifier(a.run_id or 'demo-'+uuid.uuid4().hex[:12]);out=fio.relative(a.out,rid)
 for source in [a.release,rt.CONTENT/'normalized',rt.CONTENT/'out/pack',Path('D:/Erdos_data')]:rt.disjoint_sources(source,out)
 if out.exists():raise ValueError('DEMO_RUN_EXISTS: use a new run-id')
 store=ReleaseStore(a.release);cid='cumcm';fi,refs,_=startup(store,cid,'latex','local_only');_,docrefs,_=startup(store,cid,'docx','local_only')
 task=dict(task_id=rid+'-task',competition_id=cid,is_fixture=True,title='合成示例：成本拟合与独立成本函数优化',classification=dict(mode='deterministic_fixture',llm_executed=False),subproblems=[dict(subproblem_id='q1',problem_types=['statistical_analysis','prediction'],method_tags=['linear_regression'],data_tags=['tabular']),dict(subproblem_id='q2',problem_types=['optimization'],method_tags=[],data_tags=['tabular'])])
 rt.write(out/'task.json',task);objective=dict(center=6.0,offset=3.0,lower=1.0,upper=10.0,grid_points=2001,is_fixture=True);rt.write(out/'objective.json',objective)
 import numpy as np
 rng=np.random.default_rng(42);x0=np.linspace(1,10,20);y0=2*x0+5+rng.normal(0,.6,len(x0));buf=io.StringIO();writer=csv.writer(buf,lineterminator='\n');writer.writerow(['volume','unit_cost']);writer.writerows([['%.17g'%x,'%.17g'%y] for x,y in zip(x0,y0)]);fio.write_exact(out/'data.csv',buf.getvalue().encode('utf8'))
 values=np.loadtxt(io.StringIO((out/'data.csv').read_text(encoding='utf8')),delimiter=',',skiprows=1);x,y=values[:,0],values[:,1];slope,intercept=np.polyfit(x,y,1);pred=slope*x+intercept;res=y-pred;o=rt.load(out/'objective.json');grid=np.linspace(o['lower'],o['upper'],o['grid_points']);cost=(grid-o['center'])**2+o['offset'];ix=int(np.argmin(cost));analytic=min(max(o['center'],o['lower']),o['upper'])
 results=dict(q1=dict(algorithm='OLS numpy.polyfit degree 1',input_sha256=rt.checked(out/'data.csv')['sha256'],slope=float(slope),intercept=float(intercept),r_squared=float(1-np.sum(res**2)/np.sum((y-y.mean())**2)),rmse=float(np.sqrt(np.mean(res**2))),n_samples=len(x)),q2=dict(algorithm='explicit objective grid argmin',input_sha256=rt.checked(out/'objective.json')['sha256'],center=o['center'],offset=o['offset'],opt_volume=float(grid[ix]),opt_cost=float(cost[ix]),analytic_volume=float(analytic),grid_error=float(abs(grid[ix]-analytic)),grid_points=o['grid_points']))
 rt.write(out/'results.json',results)
 _,outline=call(cmd_outline,store,SimpleNamespace(cid=cid,format='latex',model_execution='local_only',task_json=out/'task.json',budget=None));search_code,search=call(cmd_search,store,SimpleNamespace(cid=cid,tags='optimization,statistical_analysis',tags_all=False,model_execution='local_only'));rt.write(out/'outline.json',outline);rt.write(out/'task_context.json',outline['task_context'])
 import matplotlib
 matplotlib.__version__=importlib.metadata.version('matplotlib');matplotlib.use('Agg')
 import matplotlib.pyplot as plt
 matplotlib.rcParams['font.sans-serif']=['Microsoft YaHei'];matplotlib.rcParams['axes.unicode_minus']=False
 style=rt.parse(store.read(cid,refs['figure_style']['asset_id']).decode('utf8'));colors=style.get('colors') or ['#1f77b4','#d62728'];(out/'figures').mkdir(parents=True)
 fig,axs=plt.subplots(1,2,figsize=(8,3.3));axs[0].scatter(x,y,color=colors[0]);axs[0].plot(x,pred,color=colors[1]);axs[0].set(xlabel='运量 / 单位',ylabel='单位成本 / 单位',title='当前合成数据：拟合');axs[1].scatter(x,res,color=colors[0]);axs[1].axhline(0,color=colors[1]);axs[1].set(xlabel='运量 / 单位',ylabel='残差 / 单位',title='当前合成数据：残差');fig.tight_layout();fig.savefig(str(out/'figures/fit.png'),dpi=150);plt.close(fig)
 fig,ax=plt.subplots(figsize=(6,3.3));ax.plot(grid,cost,color=colors[0]);ax.scatter([grid[ix]],[cost[ix]],color=colors[1]);ax.set(xlabel='运量 / 单位',ylabel='总成本 / 单位',title='当前输入函数：成本与最优点');fig.tight_layout();fig.savefig(str(out/'figures/cost.png'),dpi=150);plt.close(fig)
 paper=out/'paper_draft';paper.mkdir();layout=store.read(cid,refs['layout']['asset_id']);rt.check_zip(layout)
 import zipfile
 with zipfile.ZipFile(io.BytesIO(layout)) as z:
  for member in z.infolist():
   if not member.is_dir():fio.write_exact(fio.relative(paper,member.filename),z.read(member))
 # Preserve the selected layout preamble and companion files; record the platform adaptation.
 original=(paper/'main.tex').read_text(encoding='utf8');preamble=original.split(r'\begin{document}',1)[0].replace('fontset=mac','fontset=windows')
 if 'graphicx' not in preamble:preamble+='\n'+r'\usepackage{graphicx}'+'\n'
 lines=[preamble,r'\begin{document}',r'\section*{合成示例（fixture）}',tex_escape(task['title'])]
 for sec in outline['dynamic_outline']:
  lines += [r'\section{'+tex_escape(sec['title'])+'}',tex_escape(narrative(sec,results))]
  if sec['section_type']=='solve' and sec['subproblem_index']:
   name='fit.png' if sec['subproblem_index']==1 else 'cost.png';lines += [r'\begin{figure}[htbp]\centering',r'\includegraphics[width=.9\textwidth]{../figures/'+name+'}',r'\caption{当前任务计算结果（合成示例）}\end{figure}']
 lines += [r'\end{document}'];(paper/'main.tex').write_bytes(('\n'.join(lines)+'\n').encode('utf8'))
 from docx import Document
 from docx.shared import Inches
 doc=Document(io.BytesIO(store.read(cid,docrefs['layout']['asset_id'])));body=doc._element.body
 for child in list(body):
  if not child.tag.endswith('}sectPr'):body.remove(child)
 title=doc.add_heading('合成示例（fixture）',0);title.paragraph_format.keep_with_next=True;title.paragraph_format.keep_together=True;doc.add_paragraph(task['title'])
 for sec in outline['dynamic_outline']:
  heading=doc.add_heading(sec['title'],1);heading.paragraph_format.keep_with_next=True;heading.paragraph_format.keep_together=True;para=doc.add_paragraph(narrative(sec,results));para.paragraph_format.keep_together=True
  if sec['section_type']=='solve' and sec['subproblem_index']:
   para.paragraph_format.keep_with_next=True;name='fit.png' if sec['subproblem_index']==1 else 'cost.png';doc.add_picture(str(out/'figures'/name),width=Inches(5.7));doc.paragraphs[-1].paragraph_format.keep_with_next=True;caption=doc.add_paragraph('图：当前任务计算结果（合成示例）');caption.paragraph_format.keep_together=True
 doc.save(str(paper/'paper.docx'))
 trace=dict(run_id=rid,release_id=store.catalog['release_id'],release_integrity_sha256=rt.checked(store.root/'integrity.json')['sha256'],competition_id=cid,family_version=fi['family_version'],ruleset_status=fi['ruleset_status'],rules_compliance='unverified',model_seam=task['classification'],model_execution='local_only',task_input_sha256=rt.checked(out/'task.json')['sha256'],data_sha256=results['q1']['input_sha256'],objective_sha256=results['q2']['input_sha256'],search_result=search,search_exit_code=search_code,layout_refs=dict(latex=refs['layout'],docx=docrefs['layout']),layout_adaptation='LaTeX preamble reused; fontset mac changed to windows if present; DOCX keeps selected styles/sections',dynamic_outline_used=True,figure_scope='current_task',historical_result_used=False,drafts=[dict(file='paper_draft/main.tex',compiled=False,note='No TeX compile performed; source and relative dependencies checked'),dict(file='paper_draft/paper.docx',rendered=False,note='Rendering evidence recorded separately after export')],tool_versions={k:importlib.metadata.version(k) for k in ['numpy','matplotlib','python-docx']})
 rt.write(out/'trace.json',trace);rt.seal(out,{'release_integrity_sha256':trace['release_integrity_sha256']},{'scope':'synthetic local demo; not LLM solving or complete paper'});rt.check_snapshot(out);print(rt.json_bytes(dict(passed=True,out=str(out),results=results)).decode('utf8'));return 0
if __name__=='__main__':
 try:sys.exit(main())
 except (OSError,ValueError,KeyError) as e:print(str(e),file=sys.stderr);sys.exit(1)
