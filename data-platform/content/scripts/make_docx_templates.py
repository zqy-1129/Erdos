"""Word 模板生成（工作包2）：CUMCM / MCM 两套 .docx 骨架。

- 输出：out/templates/{business_id}.docx（确定性字节：固定核心属性 + 归一化 zip 时间戳，
  保证同版本重建 sha256 稳定，供 content manifest 比对）。
- 版式对照各自 LaTeX 套：CUMCM（宋体 12pt / 首行缩进 2 字符 / 行距 1.72 / 三级标题 / 三线表），
  MCM（Times New Roman 12pt / Summary Sheet / 页眉 Team# + Page X of Y）。

用法：python scripts/make_docx_templates.py
"""

import sys
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

CONTENT_DIR = Path(__file__).resolve().parents[1]
OUT_DIR = CONTENT_DIR / "out" / "templates"

FIXED_DATE = datetime(2026, 1, 1, 0, 0, 0)
ZIP_FIXED_TIME = (2026, 1, 1, 0, 0, 0)


def reconfigure_stdout():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


def set_font(run, latin, ea, size, bold=None, italic=None, color="000000"):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), ea)


def add_field(paragraph, instr, placeholder):
    run = paragraph.add_run()
    for tag, attr in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        elif tag == "text":
            el = OxmlElement("w:t")
            el.text = attr
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        run._element.append(el)
    return run


def first_line_chars(style, chars=200, pt=24):
    style.paragraph_format.first_line_indent = Pt(pt)
    ind = style._element.get_or_add_pPr().find(qn("w:ind"))
    ind.set(qn("w:firstLineChars"), str(chars))


def add_toc(doc):
    para = doc.add_paragraph()
    run = add_field(para, 'TOC \\o "1-3" \\h \\z \\u', "（打开 Word 后按 F9 或右键“更新域”生成目录）")
    set_font(run, "Times New Roman", "宋体", 12)


def add_page_break(doc):
    doc.add_page_break()


def configure_normal(doc, latin, ea, size, space_multiple, indent):
    normal = doc.styles["Normal"]
    normal.font.name = latin
    normal.font.size = Pt(size)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.line_spacing = space_multiple
    normal.paragraph_format.space_after = Pt(0)
    if indent:
        first_line_chars(normal)


def configure_headings(doc, specs):
    for name, size, latin, ea, center in specs:
        style = doc.styles[name]
        style.font.name = latin
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def setup_page(doc, margins_cm):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(margins_cm)
    sec.bottom_margin = Cm(margins_cm)
    sec.left_margin = Cm(margins_cm)
    sec.right_margin = Cm(margins_cm)
    return sec


def setup_core(doc, title):
    core = doc.core_properties
    core.title = title
    core.author = "Erdos"
    core.last_modified_by = "Erdos"
    core.created = FIXED_DATE
    core.modified = FIXED_DATE
    core.revision = 1


def restart_page_number(section):
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    section._sectPr.append(pg)


def set_three_line_borders(table, keep_inside_header=True):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, val, sz in (("top", "single", "12"), ("left", "none", "0"), ("bottom", "single", "12"),
                         ("right", "none", "0"), ("insideH", "none", "0"), ("insideV", "none", "0")):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)
    if keep_inside_header:
        for cell in table.rows[0].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "0")
            bottom.set(qn("w:color"), "000000")
            tc_borders.append(bottom)
            tc_pr.append(tc_borders)


def normalize_zip(path):
    """docx 内部 zip 条目时间戳置为固定值，保证重建 sha256 稳定。"""
    with zipfile.ZipFile(path) as zin:
        names = sorted(zin.namelist())
        entries = [(name, zin.read(name)) for name in names]
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            info = zipfile.ZipInfo(name, ZIP_FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, data)
    tmp.replace(path)


def table_cell_text(table, r, c, text, latin, ea, size, bold=False):
    cell = table.cell(r, c)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_font(run, latin, ea, size, bold=bold)


# ---------------------------------------------------------------------------
# CUMCM Word 模板
# ---------------------------------------------------------------------------
FONT_CN_SERIF = "宋体"
FONT_HEI = "SimHei"
FONT_HEI_CN = "黑体"


def clear_indent(paragraph):
    """去掉从 Normal 样式继承的首行缩进。"""
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), "0")
    ind.set(qn("w:firstLine"), "0")


def build_cumcm_docx(path):
    doc = Document()
    setup_page(doc, margins_cm=2.5)
    configure_normal(doc, "Times New Roman", FONT_CN_SERIF, 12, 1.72, indent=True)
    configure_headings(doc, (
        ("Heading 1", 17, FONT_HEI, FONT_HEI_CN, True),
        ("Heading 2", 14, FONT_HEI, FONT_HEI_CN, False),
        ("Heading 3", 12, FONT_HEI, FONT_HEI_CN, False),
    ))
    setup_core(doc, "CUMCM 国赛 Word 论文模板")

    def para(text, latin="Times New Roman", ea=FONT_CN_SERIF, size=12, bold=None, align=None, indent=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent is False:
            clear_indent(p)
        r = p.add_run(text)
        set_font(r, latin, ea, size, bold=bold)
        return p

    # 第一页：标题 + 摘要（不编号）
    para("论文标题", latin=FONT_HEI, ea=FONT_HEI_CN, size=17, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para("摘要", latin=FONT_HEI, ea=FONT_HEI_CN, size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para("针对问题……本文建立了……模型，采用……算法求解，得到……结果（约 400 字，概述问题、方法与主要数值结论）。")
    kp = para("", indent=False)
    r1 = kp.add_run("关键字：")
    set_font(r1, FONT_HEI, FONT_HEI_CN, 12, bold=True)
    r2 = kp.add_run("关键词1；关键词2；关键词3")
    set_font(r2, "Times New Roman", FONT_CN_SERIF, 12)

    # 分节：目录与正文（页码从 1 开始，页脚居中）
    sec2 = doc.add_section()
    sec2.page_width = Cm(21)
    sec2.page_height = Cm(29.7)
    sec2.top_margin = sec2.bottom_margin = sec2.left_margin = sec2.right_margin = Cm(2.5)
    restart_page_number(sec2)
    sec2.footer.is_linked_to_previous = False
    footer = sec2.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(add_field(footer, "PAGE", "1"), "Times New Roman", FONT_CN_SERIF, 10.5)

    para("目录", latin=FONT_HEI, ea=FONT_HEI_CN, size=17, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    add_toc(doc)
    add_page_break(doc)

    doc.add_heading("一、问题重述", level=1)
    para("（重述题目背景、条件与待解决的问题，说明本文要完成的任务。）")
    doc.add_heading("二、问题分析", level=1)
    para("（分析问题类型、可用方法、总体思路与技术路线，可给出流程图。）")
    doc.add_heading("三、模型假设", level=1)
    para("（列出假设并逐条说明合理性。）")
    doc.add_heading("四、符号说明", level=1)
    para("表 1  符号说明", latin=FONT_HEI, ea=FONT_HEI_CN, size=10.5, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    table = doc.add_table(rows=3, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    headers = ("符号", "含义", "单位")
    rows = (("x_i", "第 i 个决策变量", "—"), ("θ", "模型参数", "—"))
    for c, text in enumerate(headers):
        table_cell_text(table, 0, c, text, FONT_HEI, FONT_HEI_CN, 10.5, bold=True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            table_cell_text(table, r, c, text, "Times New Roman", FONT_CN_SERIF, 10.5)
    set_three_line_borders(table)

    doc.add_heading("五、模型的建立与求解", level=1)
    doc.add_heading("5.1 数据预处理", level=2)
    para("（缺失值处理、标准化等，并给出处理后数据概览。）")
    para("此处插入公式（建议使用 Word 公式编辑器输入），右侧标注编号，如：  E = mc²                    (1)",
         indent=False)
    doc.add_heading("5.2 模型的建立", level=2)
    para("（给出目标函数、约束条件与求解算法；引用文献用上标 [1]。）")
    doc.add_heading("5.2.1 目标函数与约束", level=3)
    para("（三级标题示例，按需增删。）")
    doc.add_heading("六、灵敏度分析", level=1)
    para("（关键参数扰动对结果的影响分析，给出图表。）")
    doc.add_heading("七、模型的评价与推广", level=1)
    para("（优点、缺点与改进方向，推广场景。）")
    doc.add_heading("参考文献", level=1)
    para("[1] 姜启源, 谢金星, 叶俊. 数学模型[M]. 5 版. 北京: 高等教育出版社, 2018.", indent=False)
    doc.add_heading("附录 A  核心代码", level=1)
    para("print('核心求解代码粘贴于此')  # 等宽字体，可用 Consolas 10.5pt",
         latin="Consolas", ea=FONT_CN_SERIF, size=10.5, indent=False)

    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# MCM Word 模板
# ---------------------------------------------------------------------------
def build_mcm_docx(path):
    doc = Document()
    sec = setup_page(doc, margins_cm=2.54)
    configure_normal(doc, "Times New Roman", FONT_CN_SERIF, 12, 1.15, indent=False)
    configure_headings(doc, (
        ("Heading 1", 13, "Times New Roman", FONT_CN_SERIF, False),
        ("Heading 2", 12, "Times New Roman", FONT_CN_SERIF, False),
        ("Heading 3", 12, "Times New Roman", FONT_CN_SERIF, False),
    ))
    setup_core(doc, "MCM/ICM Word Paper Template")

    def para(text, size=12, bold=None, align=None, italic=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        r = p.add_run(text)
        set_font(r, "Times New Roman", FONT_CN_SERIF, size, bold=bold, italic=italic)
        return p

    # Page 1：Summary Sheet
    head = doc.add_table(rows=2, cols=3)
    head.autofit = True
    labels = ("Problem Chosen", "2026 MCM/ICM Summary Sheet", "Team Control Number")
    values = ("A", "", "0000")
    for c in range(3):
        table_cell_text(head, 0, c, labels[c], "Times New Roman", FONT_CN_SERIF, 12, bold=True)
        table_cell_text(head, 1, c, values[c], "Times New Roman", FONT_CN_SERIF, 14, bold=True)
    rule = para("")
    rule.paragraph_format.space_before = Pt(6)
    ppr = rule._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    ppr.append(pbdr)
    para("Title of Your Paper", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("")
    para("Summary", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("Write a 300-500 word summary: problem overview, methods for each sub-problem, "
         "key numerical results, and conclusions.")
    kp = para("")
    r1 = kp.add_run("Keywords: ")
    set_font(r1, "Times New Roman", FONT_CN_SERIF, 12, bold=True)
    r2 = kp.add_run("keyword1; keyword2; keyword3")
    set_font(r2, "Times New Roman", FONT_CN_SERIF, 12, italic=True)

    # Page 2 起：目录 + 正文（页眉 Team# + Page X of Y）
    sec2 = doc.add_section()
    sec2.page_width = Cm(21)
    sec2.page_height = Cm(29.7)
    sec2.top_margin = sec2.bottom_margin = sec2.left_margin = sec2.right_margin = Inches(1)
    header = sec2.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r1 = hp.add_run("Team #0000\tPage ")
    set_font(r1, "Times New Roman", FONT_CN_SERIF, 12)
    set_font(add_field(hp, "PAGE", "1"), "Times New Roman", FONT_CN_SERIF, 12)
    r2 = hp.add_run(" of ")
    set_font(r2, "Times New Roman", FONT_CN_SERIF, 12)
    set_font(add_field(hp, "NUMPAGES", "4"), "Times New Roman", FONT_CN_SERIF, 12)

    para("Contents", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_toc(doc)
    add_page_break(doc)

    sections = (
        ("1 Introduction", "Background, restatement of the problem, and our approach."),
        ("2 Assumptions and Justifications", "List each assumption with justification."),
        ("3 Model Design", "Notations, model formulation, and solution strategy."),
        ("4 Solution and Results", "Numerical results for each sub-problem with tables/figures."),
        ("5 Sensitivity Analysis", "How key parameters affect the results."),
        ("6 Strengths and Weaknesses", "Model evaluation."),
        ("7 Conclusions", "Conclusions and possible extensions."),
    )
    for index, (title, hint) in enumerate(sections):
        doc.add_heading(title, level=1)
        para(hint)
        if index == 2:
            doc.add_heading("3.1 Notation", level=2)
            para("Sample notation table and formulation.")
    doc.add_heading("References", level=1)
    para("[1] Author A, Author B. Title of the paper[J]. Journal Name, 2020, 12(3): 45-67.")
    doc.add_heading("Appendix A  Core Code", level=1)
    para("print('paste core code here')", size=10.5)

    doc.save(str(path))
    return path


def main():
    reconfigure_stdout()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = (
        ("cumcm-docx-official-v1.docx", build_cumcm_docx),
        ("mcm-docx-official-v1.docx", build_mcm_docx),
    )
    for name, builder in outputs:
        path = OUT_DIR / name
        builder(path)
        normalize_zip(path)
        print("[OK] {} | {} bytes".format(path.name, path.stat().st_size))


if __name__ == "__main__":
    main()
