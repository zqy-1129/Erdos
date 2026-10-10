"""Read-only attachment profiles; raw formula/cache semantics are preserved, never recalculated as source facts."""
import argparse,csv,io,json,re,sys
from . import core,audit

def serial(value):
    if isinstance(value,(str,int,float,bool)) or value is None:return value
    return str(value)

def profile(row,args):
    path=audit.source_path(row,args)
    if core.sha256_of(path)!=row['sha256']:raise ValueError('attachment source hash changed')
    base={'source_id':row['source_id'],'source_sha256':row['sha256'],'year':row['year'],'relative_path':row['relative_path'],
          'problem_code_candidate':audit.problem_code(row),'extension':row['extension'],'status':'indexed','sheets':[],
          'units_status':'only literal header candidates; semantic review required','original_modified':False}
    ext=row['extension']
    if ext=='.xlsx':
        import openpyxl
        book=openpyxl.load_workbook(str(path),read_only=True,data_only=False)
        try:
            for sheet in book.worksheets:
                formula_count=0;nonempty=0;preview=[]
                if sheet.max_row*sheet.max_column>10000000:raise ValueError('sheet dimension resource limit')
                for ri,values in enumerate(sheet.iter_rows(values_only=True),1):
                    nonempty+=sum(v is not None for v in values)
                    formula_count+=sum(isinstance(v,str) and v.startswith('=') for v in values)
                    if ri<=8:preview.append({'row':ri,'values':[serial(v) for v in values[:40]]})
                base['sheets'].append({'name':sheet.title,'reported_rows':sheet.max_row,'reported_columns':sheet.max_column,'nonempty_cells':nonempty,'formula_cells':formula_count,
                                       'header_candidates':preview,'locator_system':'sheet + 1-based row/column','formula_policy':'formula strings preserved in original; cached results not treated as verified calculations'})
        finally:book.close()
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(path),on_demand=True)
        try:
            for sheet in book.sheets():
                preview=[{'row':ri+1,'values':[serial(v) for v in sheet.row_values(ri)[:40]]} for ri in range(min(8,sheet.nrows))]
                base['sheets'].append({'name':sheet.name,'reported_rows':sheet.nrows,'reported_columns':sheet.ncols,'header_candidates':preview,
                                       'locator_system':'sheet + 1-based row/column','formula_policy':'xlrd exposes cached values; raw XLS is authoritative; caches are not validated computation outputs'})
        finally:book.release_resources()
    elif ext=='.csv':
        data=path.read_bytes();encoding=None
        for candidate in ('utf-8-sig','gb18030','utf-16'):
            try:text=data.decode(candidate);encoding=candidate;break
            except UnicodeError:pass
        if encoding is None:raise ValueError('CSV encoding unrecognized')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',\t;')
        except csv.Error:dialect=csv.excel
        reader=csv.reader(io.StringIO(text),dialect)
        preview=[];count=0;columns=0
        for count,values in enumerate(reader,1):
            columns=max(columns,len(values))
            if count<=8:preview.append({'row':count,'values':values[:40]})
        base['sheets'].append({'name':'csv','rows':count,'max_columns':columns,'encoding':encoding,'delimiter':dialect.delimiter,'header_candidates':preview,'locator_system':'1-based data row/column'})
    elif ext in ('.mp4','.wmv'):
        base['media_status']='original bytes indexed; no fabricated frame/time/sampling rate metadata'
    return base

def main(args=None):
    audit.deps()
    if args is None:
        args=argparse.Namespace(run_id='cumcm-2010-2025-codex-r002',data_root='D:/Erdos_data')
    rd=audit.run_dir(args)
    unique={r['sha256']:r for r in audit.all_sources(args) if r['root_kind']=='problem' and r['extension'] in ('.xlsx','.xls','.csv','.mp4','.wmv')}
    rows=[];failures=[]
    for row in unique.values():
        try:rows.append(profile(row,args))
        except Exception as exc:failures.append({'source_id':row['source_id'],'reason':str(exc),'original_indexed':True})
        if len(rows)%50==0:print('attachment profiles {}'.format(len(rows)),flush=True)
    core.write_jsonl(rd/'catalog/attachment_profiles.jsonl',rows)
    core.write_json(rd/'quality/attachment_report.json',{'input_unique_files':len(unique),'profiled':len(rows),'failures':failures,'mode':'read-only original values; no workbook exports'})
    print('attachments complete: profiled={} failures={}'.format(len(rows),len(failures)))

if __name__=='__main__':main()
